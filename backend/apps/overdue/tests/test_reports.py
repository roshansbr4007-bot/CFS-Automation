"""Phase 9: overdue-reason report (JSON / CSV / XLSX): scope, filters, columns, permissions."""

import csv
import io
from datetime import datetime

import pytest
from openpyxl import load_workbook

from apps.accounts import roles
from apps.overdue import reports
from apps.overdue.models import OverdueCase

pytestmark = pytest.mark.django_db
REPORT = "/api/v1/overdue-cases/reports/"
CSV_URL = "/api/v1/overdue-cases/reports/export/csv/"
XLSX_URL = "/api/v1/overdue-cases/reports/export/excel/"


@pytest.fixture
def hr(staff):
    user, _ = staff(roles.HR, department="HR")
    return user


@pytest.fixture
def cases(client_for, ops, work, hr):
    """Rahul: REVIEWED (late completion at 11:20, reason DEPENDENCY, cause SYSTEM).
    Amit: OPEN (still running)."""
    reviewed = work.raise_task(ops["rahul_emp"], 2026, 10, 5, 10, 0, title='Map "RM", codes')
    work.complete(reviewed, ops["rahul"], 2026, 10, 5, 11, 20)
    work.raise_task(ops["amit_emp"], 2026, 10, 6, 10, 0, title="Upload feed", priority="HIGH")
    work.tick(2026, 10, 6, 11, 0)
    case = OverdueCase.objects.get(employee=ops["rahul_emp"])
    rahul = client_for(ops["rahul"])
    rahul.post(f"/api/v1/overdue-cases/{case.pk}/submit/",
               {"version": case.version, "reason_category": "DEPENDENCY",
                "explanation": "Waiting, for the file\nfrom RTA"}, format="json")
    case.refresh_from_db()
    client_for(ops["manager"]).post(f"/api/v1/overdue-cases/{case.pk}/review/",
                                    {"version": case.version, "cause": "SYSTEM",
                                     "remark": "Portal outage"}, format="json")
    return {"reviewed": OverdueCase.objects.get(employee=ops["rahul_emp"]),
            "open": OverdueCase.objects.get(employee=ops["amit_emp"])}


def _ids(client, **params):
    return [row["id"] for row in client.get(REPORT, params).json()["results"]]


def test_report_scope(client_for, admin_client, ops, cases, hr):
    both = {cases["reviewed"].pk, cases["open"].pk}
    assert set(_ids(admin_client)) == both and set(_ids(client_for(hr))) == both
    assert set(_ids(client_for(ops["manager"]))) == both  # OPS task department
    assert _ids(client_for(ops["rahul"])) == [cases["reviewed"].pk]  # own case only


@pytest.mark.parametrize(
    ("params", "which"),
    [({"status": "REVIEWED"}, {"reviewed"}), ({"status": "OPEN"}, {"open"}),
     ({"cause": "SYSTEM"}, {"reviewed"}), ({"reason_category": "DEPENDENCY"}, {"reviewed"}),
     ({"priority": "HIGH"}, {"reviewed", "open"}), ({"date_from": "2026-10-06"}, {"open"}),
     ({"date_to": "2026-10-05"}, {"reviewed"}), ({"date_from": "2026-10-07"}, set())],
)
def test_report_filters(admin_client, cases, params, which):
    assert set(_ids(admin_client, **params)) == {cases[w].pk for w in which}


def test_employee_department_and_task_filters(admin_client, ops, cases):
    assert _ids(admin_client, employee=ops["amit_emp"].pk) == [cases["open"].pk]
    assert set(_ids(admin_client, department=cases["open"].department_id)) == {
        cases["reviewed"].pk, cases["open"].pk,
    }
    assert _ids(admin_client, task=cases["reviewed"].task_id) == [cases["reviewed"].pk]
    for bad in ({"status": "X"}, {"cause": "X"}, {"date_from": "soon"}, {"employee": "x"},
                {"date_from": "2026-10-07", "date_to": "2026-10-01"}):
        assert admin_client.get(REPORT, bad).status_code == 400


def test_exports_are_for_hr_and_admin_only(client_for, admin_client, ops, cases, hr):
    for url in (CSV_URL, XLSX_URL):
        assert client_for(ops["rahul"]).get(url).status_code == 403
        assert client_for(ops["manager"]).get(url).status_code == 403
        assert client_for(hr).get(url).status_code == 200
        assert admin_client.get(url).status_code == 200


def test_csv_columns_values_and_escaping(admin_client, cases):
    response = admin_client.get(CSV_URL)
    assert response["Content-Type"] == "text/csv; charset=utf-8"
    assert response["Content-Disposition"] == 'attachment; filename="overdue-cases.csv"'
    rows = list(csv.reader(io.StringIO(response.content.decode("utf-8-sig"))))
    assert tuple(rows[0]) == reports.COLUMNS and len(rows) == 3
    reviewed = next(r for r in rows[1:] if r[0] == str(cases["reviewed"].pk))
    data = dict(zip(reports.COLUMNS, reviewed, strict=True))
    assert data["Task Title"] == 'Map "RM", codes'  # quotes and commas survive
    assert data["Employee Explanation"] == "Waiting, for the file\nfrom RTA"
    assert (data["Employee Reason"], data["Authoritative Cause"]) == ("DEPENDENCY", "SYSTEM")
    assert data["Reviewer Remark"] == "Portal outage" and data["Status"] == "REVIEWED"
    assert (data["Deadline"], data["Overdue At"]) == ("2026-10-05 11:00", "2026-10-05 11:00")
    assert data["Completed At"] == "2026-10-05 11:20" and data["Overdue Minutes"] == "20"
    empty = admin_client.get(CSV_URL, {"status": "REVIEWED", "cause": "CLIENT"})
    assert list(csv.reader(io.StringIO(empty.content.decode("utf-8-sig")))) == [
        list(reports.COLUMNS)
    ]


def test_excel_structure_and_values(admin_client, cases):
    response = admin_client.get(XLSX_URL)
    assert response["Content-Type"] == reports.XLSX_CONTENT_TYPE
    book = load_workbook(io.BytesIO(response.content))
    assert book.sheetnames == ["Overdue Cases"]
    sheet = book["Overdue Cases"]
    assert [c.value for c in sheet[1]] == list(reports.COLUMNS) and sheet.max_row == 3
    assert all(c.font.bold for c in sheet[1]) and sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == "A1:U3"
    row = next(r for r in sheet.iter_rows(min_row=2) if r[0].value == cases["reviewed"].pk)
    values = dict(zip(reports.COLUMNS, (c.value for c in row), strict=True))
    assert values["Deadline"] == datetime(2026, 10, 5, 11, 0)  # IST, without a time zone
    assert row[9].number_format == "yyyy-mm-dd hh:mm"
    assert (values["Employee Reason"], values["Authoritative Cause"]) == ("DEPENDENCY", "SYSTEM")
    assert values["Overdue Minutes"] == 20
    filtered = load_workbook(io.BytesIO(admin_client.get(XLSX_URL, {"status": "OPEN"}).content))
    assert filtered["Overdue Cases"].max_row == 2
    empty = load_workbook(io.BytesIO(admin_client.get(XLSX_URL, {"cause": "CLIENT"}).content))
    assert empty["Overdue Cases"].max_row == 1
