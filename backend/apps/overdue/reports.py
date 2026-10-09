"""Overdue-reason report rows and CSV / XLSX exports (Phase 9). Read-only.

Follows the Phase 7 export conventions (UTF-8 CSV with a byte-order mark; an in-memory openpyxl
workbook with a bold frozen header, autofilter and readable widths) without reusing or changing
Phase 7 code. Times are IST. The overdue duration runs to completion (or now, while open)."""

import csv
import io

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from apps.core.timeutils import to_ist

COLUMNS = (
    "Case ID", "Task Reference", "Task Title", "Employee ID", "Employee Name", "Department",
    "Priority", "Category", "SLA Start", "Deadline", "Overdue At", "Completed At",
    "Overdue Minutes", "Status", "Employee Reason", "Employee Explanation", "Submitted At",
    "Reviewer", "Authoritative Cause", "Reviewer Remark", "Reviewed At",
)
DATETIME_COLUMNS = {9, 10, 11, 12, 17, 21}  # 1-based column numbers holding times
XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
REPORT_ORDER = ("department__code", "employee__full_name", "employee_id", "opened_at", "id")


def overdue_minutes(case, now=None) -> int:
    """From the overdue moment to completion (frozen) or to now while the task is open."""
    end = case.clock.stopped_at or now or timezone.now()
    return max(0, int((end - case.overdue_at).total_seconds() // 60))


def _who(user) -> str:
    if user is None:
        return ""
    employee = getattr(user, "employee", None)
    return employee.full_name if employee is not None else user.email


def row(case, now=None) -> tuple:
    employee = case.employee
    return (
        case.pk, case.task.reference, case.task_title,
        employee.employee_code or str(employee.pk), employee.full_name, case.department.code,
        case.priority, case.category_name, case.sla_start_at, case.sla_due_at, case.overdue_at,
        case.task.completed_at, overdue_minutes(case, now), case.status, case.reason_category,
        case.explanation, case.submitted_at, _who(case.reviewed_by), case.cause,
        case.review_remark, case.reviewed_at,
    )


def _csv_value(value) -> str:
    if value is None:
        return ""
    if hasattr(value, "tzinfo"):
        return to_ist(value).strftime("%Y-%m-%d %H:%M")
    return str(value)


def write_csv(cases, now=None) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\r\n")
    writer.writerow(COLUMNS)
    for case in cases:
        writer.writerow([_csv_value(v) for v in row(case, now)])
    return buffer.getvalue().encode("utf-8-sig")


def _cell(value):
    if value is None or value == "":
        return None
    if hasattr(value, "tzinfo"):  # Excel cannot store time zones: write naive IST times
        return to_ist(value).replace(tzinfo=None)
    return value


def write_xlsx(cases, now=None) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Overdue Cases"
    sheet.append(list(COLUMNS))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    count = 0
    for case in cases:
        sheet.append([_cell(v) for v in row(case, now)])
        count += 1
    for index, name in enumerate(COLUMNS, start=1):
        letter = get_column_letter(index)
        sheet.column_dimensions[letter].width = max(12, len(name) + 4)
        if index in DATETIME_COLUMNS:
            for (cell,) in sheet.iter_rows(min_row=2, min_col=index, max_col=index):
                cell.number_format = "yyyy-mm-dd hh:mm"
    for letter, width in (("C", 36), ("P", 48), ("T", 48)):
        sheet.column_dimensions[letter].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{count + 1}"
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
