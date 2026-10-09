"""Phase 7 Stage B: performance report, CSV and Excel export (read-only).

Data is produced only through the frozen Stage A services. Rahul (OPS) has a FINALIZED
October (95, Excellent); Amit (OPS) has a CALCULATED October; an RM employee has FINALIZED
September (85, Very Good) and October (75, Good). With no task work, Timeliness is not
applicable, so it is scored through the approved Option B override."""

import csv
import io

import pytest
from openpyxl import load_workbook

from apps.accounts import roles
from apps.audit.models import AuditLog
from apps.performance import reports, services
from apps.performance.models import MonthlyPerformance

pytestmark = pytest.mark.django_db
REPORT = "/api/v1/performance/reports/"
CSV_URL = "/api/v1/performance/reports/export/csv/"
XLSX_URL = "/api/v1/performance/reports/export/excel/"
KPI_CODES = ["ACCURACY", "CLIENT_SERVICING", "COMPLIANCE", "DATA_SYSTEM", "FINANCIAL_ACCURACY",
             "TIMELINESS"]
TRICKY_REMARK = 'Strong month, "clear" ownership,\nwell done'


def _finalized(actor, employee, year, month, score, remark=""):
    performance = services.calculate_monthly_performance(actor=actor, employee=employee,
                                                         year=year, month=month)
    for code in KPI_CODES:
        performance.refresh_from_db()
        services.update_kpi_score(actor=actor, performance=performance,
                                  version=performance.version, kpi_code=code, score=score)
    performance.refresh_from_db()
    performance = services.submit_review(actor=actor, performance=performance,
                                         version=performance.version, remark=remark)
    return services.finalize_performance(actor=actor, performance=performance,
                                         version=performance.version)


@pytest.fixture
def data(admin_user, ops, staff, assign):
    rm_user, rm_emp = staff(roles.EMPLOYEE, department="RM")
    for employee in (ops["rahul_emp"], ops["amit_emp"], rm_emp):
        assign(employee)
    _finalized(admin_user, ops["rahul_emp"], 2026, 10, "95", TRICKY_REMARK)
    services.calculate_monthly_performance(actor=admin_user, employee=ops["amit_emp"],
                                           year=2026, month=10)
    _finalized(admin_user, rm_emp, 2026, 9, "85")
    _finalized(admin_user, rm_emp, 2026, 10, "75")
    return {"rm_user": rm_user, "rm_emp": rm_emp}


@pytest.fixture
def hr_client(client_for, staff):
    hr, _ = staff(roles.HR, department="HR")
    return client_for(hr)


def _rows(client, **params):
    body = client.get(REPORT, params).json()
    return body["results"]


def _key(row):
    return (row["employee"]["department"]["code"], row["employee"]["full_name"],
            row["employee"]["id"], row["year"], row["month"])


# --- report -----------------------------------------------------------------------------------


def test_report_returns_the_stored_snapshot(admin_client, ops, data):
    body = admin_client.get(REPORT).json()
    assert body["count"] == 4
    rahul = next(r for r in body["results"] if r["employee"]["id"] == ops["rahul_emp"].pk)
    assert (rahul["status"], rahul["overall_score"], rahul["performance_band"]) == (
        "FINALIZED", "95.00", "Excellent",
    )
    assert rahul["manager_remark"] == TRICKY_REMARK
    assert rahul["employee"]["employee_id"] == reports.employee_id(ops["rahul_emp"])
    assert rahul["employee"]["department"]["code"] == "OPS"
    assert [k["code"] for k in rahul["kpi_scores"]] == KPI_CODES  # KPI-code order
    accuracy = rahul["kpi_scores"][0]
    assert accuracy == {"code": "ACCURACY", "name": "Accuracy", "weight": "3.00",
                        "score": "95.00", "source": "MANAGER", "remarks": ""}
    timeliness = rahul["kpi_scores"][-1]
    assert (timeliness["weight"], timeliness["source"]) == ("2.00", "MANAGER")  # Option B
    amit = next(r for r in body["results"] if r["employee"]["id"] == ops["amit_emp"].pk)
    assert (amit["status"], amit["overall_score"], amit["performance_band"]) == (
        "CALCULATED", None, "",
    )


def test_order_is_department_name_year_month(admin_client, data):
    rows = _rows(admin_client)
    assert [_key(r) for r in rows] == sorted(_key(r) for r in rows)
    assert [r["employee"]["department"]["code"] for r in rows] == ["OPS", "OPS", "RM", "RM"]
    assert [(r["year"], r["month"]) for r in rows[2:]] == [(2026, 9), (2026, 10)]


def test_empty_result(admin_client, data):
    body = admin_client.get(REPORT, {"year": 2030}).json()
    assert body["count"] == 0 and body["results"] == []


@pytest.mark.parametrize(
    ("params", "expected"),
    [
        ({"year": 2026}, 4),
        ({"month": 9}, 1),
        ({"status": "CALCULATED"}, 1),
        ({"status": "FINALIZED"}, 3),
        ({"performance_band": "Good"}, 1),
        ({"performance_band": "Excellent"}, 1),
        ({"date_from": "2026-10-01", "date_to": "2026-10-31"}, 3),
        ({"date_to": "2026-09-30"}, 1),
        ({"date_from": "2026-11-01"}, 0),
    ],
)
def test_simple_filters(admin_client, data, params, expected):
    assert len(_rows(admin_client, **params)) == expected


def test_employee_department_kpi_and_combined_filters(admin_client, ops, data):
    assert [r["employee"]["id"] for r in _rows(admin_client, employee=ops["rahul_emp"].pk)] == [
        ops["rahul_emp"].pk
    ]
    rm = data["rm_emp"]
    assert {r["employee"]["id"] for r in _rows(admin_client, department=rm.department_id)} == {
        rm.pk
    }
    accuracy_only = _rows(admin_client, kpi="ACCURACY")
    assert len(accuracy_only) == 4
    assert {tuple(k["code"] for k in r["kpi_scores"]) for r in accuracy_only} == {("ACCURACY",)}
    combined = _rows(admin_client, department=rm.department_id, month=10,
                     status="FINALIZED", performance_band="Good")
    assert [(r["employee"]["id"], r["month"]) for r in combined] == [(rm.pk, 10)]


@pytest.mark.parametrize(
    "params",
    [{"month": "13"}, {"year": "abc"}, {"status": "DONE"}, {"performance_band": "Great"},
     {"date_from": "yesterday"}, {"kpi": "NOPE"}, {"employee": "x"},
     {"date_from": "2026-10-31", "date_to": "2026-10-01"}],
)
def test_invalid_filters_are_rejected(admin_client, params):
    assert admin_client.get(REPORT, params).status_code == 400


# --- permissions ------------------------------------------------------------------------------


def test_anonymous_is_rejected(api_client):
    for url in (REPORT, CSV_URL, XLSX_URL):
        assert api_client.get(url).status_code == 401


def test_employee_sees_only_own_records_in_any_status(admin_user, client_for, ops, data):
    """Employee -> own records only, whatever their status (no finalized-only restriction)."""
    services.calculate_monthly_performance(actor=admin_user, employee=ops["amit_emp"],
                                           year=2026, month=9)
    september = MonthlyPerformance.objects.get(employee=ops["amit_emp"], year=2026, month=9)
    services.submit_review(actor=admin_user, performance=september, version=september.version)
    amit, rahul = client_for(ops["amit"]), client_for(ops["rahul"])
    assert {(r["employee"]["id"], r["status"]) for r in _rows(amit)} == {
        (ops["amit_emp"].pk, "CALCULATED"), (ops["amit_emp"].pk, "UNDER_REVIEW"),
    }
    assert {(r["employee"]["id"], r["status"]) for r in _rows(rahul)} == {
        (ops["rahul_emp"].pk, "FINALIZED"),
    }
    # Another employee's records are never visible, however the filters are set.
    assert _rows(rahul, employee=ops["amit_emp"].pk) == []
    assert _rows(amit, employee=ops["rahul_emp"].pk) == []
    assert _rows(rahul, department=data["rm_emp"].department_id) == []
    assert _rows(amit, department=ops["amit_emp"].department_id, status="FINALIZED") == []
    # A department filter only narrows the employee's own scope.
    assert len(_rows(amit, department=ops["amit_emp"].department_id)) == 2


def test_manager_sees_own_department_only(client_for, ops, data):
    manager = client_for(ops["manager"])
    assert {r["employee"]["id"] for r in _rows(manager)} == {
        ops["rahul_emp"].pk, ops["amit_emp"].pk,
    }
    assert _rows(manager, department=data["rm_emp"].department_id) == []
    assert _rows(manager, employee=data["rm_emp"].pk) == []


def test_hr_and_admin_see_everything(admin_client, hr_client, data):
    assert len(_rows(hr_client)) == len(_rows(admin_client)) == 4


def test_only_hr_and_admin_can_export(admin_client, hr_client, client_for, ops, data):
    for url in (CSV_URL, XLSX_URL):
        assert client_for(ops["rahul"]).get(url).status_code == 403
        assert client_for(ops["manager"]).get(url).status_code == 403
        assert hr_client.get(url).status_code == 200
        assert admin_client.get(url).status_code == 200


# --- CSV --------------------------------------------------------------------------------------


def _csv(client, **params):
    response = client.get(CSV_URL, params)
    assert response.status_code == 200
    return response, list(csv.reader(io.StringIO(response.content.decode("utf-8-sig"))))


def test_csv_structure_order_and_escaping(admin_client, ops, data):
    response, rows = _csv(admin_client)
    assert response["Content-Type"] == "text/csv; charset=utf-8"
    assert response["Content-Disposition"] == 'attachment; filename="performance-report.csv"'
    assert response.content.startswith("\ufeff".encode())  # UTF-8 byte-order mark
    assert tuple(rows[0]) == reports.CSV_COLUMNS
    body = rows[1:]
    assert len(body) == 4 * 6  # one row per employee-month and KPI
    keys = [(r[0], r[3], r[4], r[9]) for r in body]
    assert len(set(keys)) == len(keys)  # no duplicate employee/KPI rows
    report_order = [(r["employee"]["employee_id"], str(r["year"]), str(r["month"]))
                    for r in _rows(admin_client)]
    assert list(dict.fromkeys((r[0], r[3], r[4]) for r in body)) == report_order
    rahul = [r for r in body if r[0] == reports.employee_id(ops["rahul_emp"])]
    assert [r[9] for r in rahul] == KPI_CODES
    assert rahul[0][6:9] == ["95.00", "Excellent", TRICKY_REMARK]  # escaped and round-trips
    assert rahul[0][11:14] == ["3.00", "95.00", "MANAGER"]


def test_csv_follows_filters_and_handles_empty(admin_client, data):
    _, rm_rows = _csv(admin_client, department=data["rm_emp"].department_id)
    assert len(rm_rows) == 1 + 2 * 6
    assert {r[2] for r in rm_rows[1:]} == {"RM"}
    _, accuracy = _csv(admin_client, kpi="ACCURACY")
    assert {r[9] for r in accuracy[1:]} == {"ACCURACY"} and len(accuracy) == 1 + 4
    _, empty = _csv(admin_client, year=2030)
    assert empty == [list(reports.CSV_COLUMNS)]  # headers, no data rows


# --- Excel ------------------------------------------------------------------------------------


def _workbook(client, **params):
    response = client.get(XLSX_URL, params)
    assert response.status_code == 200
    return response, load_workbook(io.BytesIO(response.content))


def test_excel_structure_and_values(admin_client, ops, data):
    response, book = _workbook(admin_client)
    assert response["Content-Type"] == reports.XLSX_CONTENT_TYPE
    assert response["Content-Disposition"] == 'attachment; filename="performance-report.xlsx"'
    assert book.sheetnames == ["Performance Summary", "KPI Details"]
    summary, details = book["Performance Summary"], book["KPI Details"]
    assert [c.value for c in summary[1]] == list(reports.SUMMARY_COLUMNS)
    assert [c.value for c in details[1]] == list(reports.KPI_COLUMNS)
    assert (summary.max_row, details.max_row) == (1 + 4, 1 + 4 * 6)
    assert all(c.font.bold for c in summary[1]) and all(c.font.bold for c in details[1])
    assert summary.freeze_panes == "A2" and details.freeze_panes == "A2"
    assert summary.auto_filter.ref == "A1:H5" and details.auto_filter.ref == "A1:J25"
    rahul = next(row for row in summary.iter_rows(min_row=2, values_only=True)
                 if row[0] == reports.employee_id(ops["rahul_emp"]))
    assert rahul[2:] == ("OPS", 2026, 10, "FINALIZED", 95.0, "Excellent")
    overall = next(c for c in summary["G"][1:] if c.value == 95.0)
    assert overall.number_format == "0.00"
    accuracy = next(row for row in details.iter_rows(min_row=2)
                    if row[0].value == reports.employee_id(ops["rahul_emp"])
                    and row[4].value == "ACCURACY")
    assert [c.value for c in accuracy[4:9]] == ["ACCURACY", "Accuracy", 3.0, 95.0, "MANAGER"]
    assert accuracy[6].number_format == "0.00" and accuracy[7].number_format == "0.00"
    amit = next(row for row in summary.iter_rows(min_row=2, values_only=True)
                if row[0] == reports.employee_id(ops["amit_emp"]))
    assert amit[5:] == ("CALCULATED", None, None)  # not finalized: no score, no band


def test_excel_follows_filters_and_handles_empty(admin_client, data):
    _, book = _workbook(admin_client, department=data["rm_emp"].department_id)
    assert book["Performance Summary"].max_row == 1 + 2
    assert book["KPI Details"].max_row == 1 + 2 * 6
    _, empty = _workbook(admin_client, year=2030)
    assert empty["Performance Summary"].max_row == 1
    assert empty["KPI Details"].max_row == 1
    assert [c.value for c in empty["KPI Details"][1]] == list(reports.KPI_COLUMNS)


# --- read-only --------------------------------------------------------------------------------


def test_reading_and_exporting_change_nothing(admin_client, hr_client, data):
    before = list(MonthlyPerformance.objects.order_by("id").values_list("id", "version",
                                                                         "updated_at"))
    audit_rows = AuditLog.objects.count()
    for client in (admin_client, hr_client):
        client.get(REPORT)
        client.get(CSV_URL)
        client.get(XLSX_URL)
    after = list(MonthlyPerformance.objects.order_by("id").values_list("id", "version",
                                                                        "updated_at"))
    assert after == before
    assert AuditLog.objects.count() == audit_rows
