"""Phase 7.4 HR / Admin KRA month listing and CSV / Excel exports: one validated filter set for
all three, the department recorded at calculation, the provisional rule, all matching rows in
the files, restricted access, nothing internal in the files, and nothing changed by reading.
Also: the finalized legacy history Home reads from the UNCHANGED legacy report."""

import csv
import io
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from apps.accounts import roles
from apps.audit.models import AuditLog
from apps.org.models import Department
from apps.performance import kra_reports, services
from apps.performance import review_services as review
from apps.performance.models import MonthlyPerformance

from . import test_kra_review as shared
from .test_kra_review import _component, _deduct, _fresh, _kpi_id, _submitted

pytestmark = pytest.mark.django_db
D = Decimal
BASE = "/api/v1/performance/months"

# Shared pytest fixtures (the closed November month of test_kra_review and its HR user).
hr = shared.hr
month = shared.month


def _kra(employee, month_no, *, status="FINALIZED", total="7", band="Consistent Performer",
         cutoff=None, department=None, year=2026):
    start, end = services.month_bounds(year, month_no)
    return MonthlyPerformance.objects.create(
        employee=employee, year=year, month=month_no, period_start=start, period_end=end,
        status=status, calculation_model="KRA_POINTS",
        department=department or employee.department, final_total=D(total),
        auto_total=D(total), adjustment_total=D("0"), deduction_total=D("0"),
        max_points_applicable=D("10"), band_name=band,
        cutoff_at=cutoff or services._period_window(start, end)[1],
    )


@pytest.fixture
def year(ops):
    """36 KRA months (3 employees x 12) and one legacy month that must never appear."""
    rows = []
    for emp in (ops["rahul_emp"], ops["amit_emp"], ops["manager_emp"]):
        for month_no in range(1, 13):
            rows.append(_kra(emp, month_no, total=str(5 + month_no % 5),
                             band="Consistent Performer" if month_no % 2 else "High Performer"))
    start, end = services.month_bounds(2025, 12)
    MonthlyPerformance.objects.create(employee=ops["rahul_emp"], year=2025, month=12,
                                      period_start=start, period_end=end, status="FINALIZED")
    return rows


def _keys(rows):
    return [(r["employee"]["id"], r["year"], r["month"]) for r in rows]


def _all_list_rows(client, **params):
    rows, page = [], 1
    while True:
        body = client.get(f"{BASE}/", {**params, "page": page}).json()
        rows += body["results"]
        if not body["next"]:
            return rows
        page += 1


def _csv_rows(response):
    assert response.content.startswith(b"\xef\xbb\xbf")
    return list(csv.reader(io.StringIO(response.content.decode("utf-8-sig"))))


def _book(response):
    return load_workbook(io.BytesIO(response.content))


# --- one filter set: list, CSV and Excel agree, and the files hold every row -------------------


@pytest.mark.parametrize("params", [
    {}, {"month": 3}, {"band": "High Performer"}, {"status": "FINALIZED", "year": 2026},
])
def test_list_csv_and_excel_contain_the_same_rows(year, admin_client, params):
    rows = _all_list_rows(admin_client, **params)
    listed = {(r["employee"]["full_name"], r["year"], r["month"]) for r in rows}
    first = admin_client.get(f"{BASE}/", params).json()
    assert first["count"] == len(rows) and len(first["results"]) == min(25, len(rows))
    csv_rows = _csv_rows(admin_client.get(f"{BASE}/export/csv/", params))
    header = csv_rows[0]
    assert header == list(kra_reports.CSV_COLUMNS)
    in_csv = {(r[1], int(r[3]), int(r[4])) for r in csv_rows[1:]}
    sheet = _book(admin_client.get(f"{BASE}/export/excel/", params))["KRA Summary"]
    in_excel = {(r[1], r[3], r[4]) for r in sheet.iter_rows(min_row=2, values_only=True)}
    assert listed == in_csv == in_excel
    if not params:
        assert len(rows) == 36  # beyond one page; the legacy month never appears


def test_order_and_values(year, admin_client, ops):
    rows = _all_list_rows(admin_client)
    names = [r["employee"]["full_name"] for r in rows]
    assert names == sorted(names)  # same department: by name, then year, month
    first = rows[0]
    assert set(first) == {
        "id", "employee", "department", "year", "month", "status", "provisional", "reopen_count",
        "max_points_applicable", "auto_total", "adjustment_total", "deduction_total",
        "final_total", "band", "band_ceiling", "finalized_at",
    }
    assert first["department"]["code"] == "OPS" and first["provisional"] is False


def test_the_department_recorded_at_calculation_is_used(year, admin_client, ops):
    rm = Department.objects.get(code="RM")
    ops_id = ops["rahul_emp"].department_id
    type(ops["rahul_emp"]).objects.filter(pk=ops["rahul_emp"].pk).update(department=rm)
    in_ops = _all_list_rows(admin_client, department=ops_id, employee=ops["rahul_emp"].pk)
    assert len(in_ops) == 12 and {r["department"]["code"] for r in in_ops} == {"OPS"}
    assert _all_list_rows(admin_client, department=rm.pk) == []
    csv_rows = _csv_rows(admin_client.get(f"{BASE}/export/csv/",
                                          {"employee": ops["rahul_emp"].pk}))
    assert {r[2] for r in csv_rows[1:]} == {"OPS"}


def test_provisional_follows_the_month_closed_rule(ops, admin_client, ist):
    emp = ops["rahul_emp"]
    cases = {
        1: ist(2026, 1, 20, 9, 0),  # mid-month: provisional
        2: ist(2026, 2, 28, 23, 59),  # last minute of the month: provisional
        3: ist(2026, 4, 1, 0, 0),  # exactly the month end: closed
    }
    records = {m: _kra(emp, m, status="CALCULATED", cutoff=cutoff)
               for m, cutoff in cases.items()}
    listed = {r["month"]: r["provisional"] for r in _all_list_rows(admin_client)}
    for month_no, record in records.items():
        assert listed[month_no] is (not review.month_closed(record))
    assert listed == {1: True, 2: True, 3: False}
    assert [r["month"] for r in _all_list_rows(admin_client, provisional="true")] == [1, 2]
    assert [r["month"] for r in _all_list_rows(admin_client, provisional="false")] == [3]


@pytest.mark.parametrize("bad", [
    {"year": "abc"}, {"month": 13}, {"employee": 0}, {"department": "x"},
    {"status": "NOPE"}, {"band": "Superstar"}, {"band": "a\x00b"}, {"band": "x" * 61},
    {"provisional": "maybe"},
])
@pytest.mark.parametrize("path", ["/", "/export/csv/", "/export/excel/"])
def test_invalid_filters_are_rejected_everywhere(admin_client, bad, path):
    response = admin_client.get(f"{BASE}{path}", bad)
    assert response.status_code == 400 and set(response.json()["fields"]) == set(bad)


# --- access ------------------------------------------------------------------------------------


@pytest.mark.parametrize("path", ["/", "/export/csv/", "/export/excel/"])
def test_only_kra_readers_list_and_export(year, ops, admin_client, client_for, make_user, path):
    for user in (ops["rahul"], ops["manager"]):  # Employee; Ops Manager (P11)
        assert client_for(user).get(f"{BASE}{path}").status_code == 403
    assert client_for(make_user(roles.HR)).get(f"{BASE}{path}").status_code == 200
    assert admin_client.get(f"{BASE}{path}").status_code == 200


# --- what the files contain -----------------------------------------------------------------


def test_exports_hold_stored_points_and_recorded_deductions_but_no_review_detail(
    month, hr, admin_client
):
    plan, _, record = month
    record = _submitted(record, hr)
    record = review.adjust_kpi(actor=hr, performance=record, version=record.version,
                               kpi_id=_kpi_id(record, "ACCURACY"), points=D("2.5"),
                               reason="SECRET ADJUSTMENT REASON")
    record = _deduct(record, hr, plan, "MISSED_RECONCILIATION", D("20"),
                     component_id=_component(record, "ACCURACY").component_id)
    record = _deduct(record, hr, plan, "DATA_INCONSISTENCY", D("20"))
    record = _deduct(record, hr, plan, "TRANSACTION_ERROR")
    overall = record.deduction_applications.get(rule__code="DATA_INCONSISTENCY")
    record = review.reverse_deduction(actor=hr, performance=record,
                                      version=_fresh(record).version,
                                      application_id=overall.pk, reason="SECRET REVERSAL")
    record = _fresh(record)
    book = _book(admin_client.get(f"{BASE}/export/excel/"))
    assert book.sheetnames == ["KRA Summary", "KPI Details", "Deduction Records", "About"]
    (summary,) = book["KRA Summary"].iter_rows(min_row=2, values_only=True)
    assert summary[2] == "OPS" and summary[11] == float(record.final_total) == 6.8
    assert summary[12] == "Needs Improvement" and summary[13] == "Needs Improvement"
    kpis = {row[4]: row for row in book["KPI Details"].iter_rows(min_row=2, values_only=True)}
    accuracy = kpis["ACCURACY"]
    stored = record.kpi_scores.get(kpi__code="ACCURACY")
    assert (accuracy[10], accuracy[11], accuracy[12], accuracy[13], accuracy[14]) == (
        3.0, -0.5, float(stored.deduction_points), 0.5, float(stored.final_points),
    )
    records = [row for row in book["Deduction Records"].iter_rows(min_row=2, values_only=True)]
    assert [(r[6], r[9], r[12], r[13], r[14]) for r in records] == [
        ("MISSED_RECONCILIATION", "COMPONENT", 20.0, None, "Active"),
        ("DATA_INCONSISTENCY", "OVERALL", 20.0, None, "Reversed"),
        ("TRANSACTION_ERROR", "OVERALL", None, "Needs Improvement", "Active"),
        ("DATA_INCONSISTENCY", "OVERALL", 20.0, None, "Reversal"),
    ]
    assert "Points" not in " ".join(kra_reports.DEDUCTION_COLUMNS)
    about = " ".join(str(c) for row in book["About"].iter_rows(values_only=True) for c in row)
    assert "NOT a per-rule breakdown of points" in about
    every_cell = " ".join(str(c) for ws in book for row in ws.iter_rows(values_only=True)
                          for c in row if c is not None)
    csv_text = admin_client.get(f"{BASE}/export/csv/").content.decode("utf-8-sig")
    for secret in ("SECRET", "Ticket 42", "found in review", hr.email):
        assert secret not in every_cell and secret not in csv_text
    assert len(_csv_rows(admin_client.get(f"{BASE}/export/csv/"))) == 1 + 6  # one per KPI


def test_empty_files_still_have_headers(admin_client):
    assert _csv_rows(admin_client.get(f"{BASE}/export/csv/")) == [list(kra_reports.CSV_COLUMNS)]
    book = _book(admin_client.get(f"{BASE}/export/excel/"))
    assert [c.value for c in book["KRA Summary"][1]] == list(kra_reports.SUMMARY_COLUMNS)
    assert book["KRA Summary"].max_row == 1


def test_listing_and_exporting_change_nothing(month, admin_client, client_for, hr):
    hr_client = client_for(hr)  # signed in before the snapshot (a sign-in is audited)
    before = (list(MonthlyPerformance.objects.values_list("pk", "version", "status",
                                                           "final_total")),
              AuditLog.objects.count())
    for path in ("/", "/export/csv/", "/export/excel/"):
        assert hr_client.get(f"{BASE}{path}").status_code == 200
        assert admin_client.get(f"{BASE}{path}").status_code == 200
    after = (list(MonthlyPerformance.objects.values_list("pk", "version", "status",
                                                          "final_total")),
             AuditLog.objects.count())
    assert after == before


# --- the finalized legacy history Home reads (legacy report, unchanged) ------------------------


def test_the_legacy_report_gives_each_caller_their_own_finalized_legacy_months(
    ops, staff, client_for
):
    hr_user, hr_emp = staff(roles.HR)
    people = {"rahul": (ops["rahul"], ops["rahul_emp"]), "manager": (ops["manager"],
              ops["manager_emp"]), "hr": (hr_user, hr_emp)}
    for _, emp in people.values():
        for month_no, status in ((7, "FINALIZED"), (8, "CALCULATED")):
            start, end = services.month_bounds(2026, month_no)
            MonthlyPerformance.objects.create(employee=emp, year=2026, month=month_no,
                                              period_start=start, period_end=end,
                                              status=status)
    _kra(ops["rahul_emp"], 9)  # a KRA month never appears in the legacy report
    for user, emp in people.values():
        body = client_for(user).get("/api/v1/performance/reports/",
                                    {"employee": emp.pk, "status": "FINALIZED",
                                     "page_size": 200}).json()
        assert [(r["employee"]["id"], r["month"], r["status"]) for r in body["results"]] == [
            (emp.pk, 7, "FINALIZED"),
        ]


# --- existing contracts stay as they were ------------------------------------------------------


def test_existing_report_and_review_contracts_are_unchanged():
    from drf_spectacular.generators import SchemaGenerator

    schema = SchemaGenerator().get_schema(request=None, public=True)
    components = schema["components"]["schemas"]
    assert sorted(components["PerformanceReportRow"]["properties"]) == [
        "assigned_tasks", "completed_tasks", "completion_rate", "employee", "id", "kpi_scores",
        "manager_remark", "month", "overall_score", "overdue_tasks", "pending_tasks",
        "performance_band", "period_end", "period_start", "sla_breached_tasks",
        "sla_compliance_rate", "status", "year",
    ]
    assert sorted(components["KraRevMonth"]["properties"]) == [
        "adjustment_total", "auto_total", "band", "band_ceiling", "blockers", "cutoff_at",
        "deduction_applications", "deduction_lines", "deduction_total", "employee",
        "final_total", "finalized_at", "finalized_by", "id", "kpis", "max_points_applicable",
        "month", "plan_version", "provisional", "reopen_count", "reviewed_at", "reviewed_by",
        "status", "version", "year",
    ]
    for path in ("/api/v1/performance/reports/", "/api/v1/performance/months/{id}/"):
        assert sorted(schema["paths"][path]) == ["get"]
