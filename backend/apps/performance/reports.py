"""Performance report rows and exports (Phase 7 Stage B). Read-only.

Every value is read from the stored monthly snapshot (MonthlyPerformance / MonthlyKPIScore);
nothing is recalculated. JSON, CSV and Excel all take the same queryset from
selectors.report_queryset, so an export always matches the filtered report. Files are built
in memory and returned as HTTP responses; nothing is written to disk.

Known limitation: the Department column is the employee's CURRENT department; Stage A does not
persist a department snapshot on MonthlyPerformance, so historical months are not reported with
the department the employee had at the time.
"""

import csv
import io
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

SUMMARY_COLUMNS = (
    "Employee ID", "Employee Name", "Department", "Year", "Month", "Status",
    "Overall Score", "Performance Band",
)
KPI_COLUMNS = (
    "Employee ID", "Employee Name", "Year", "Month", "KPI Code", "KPI Name", "Weight",
    "Score", "Source", "Remarks",
)
# CSV: one row per (employee-month, KPI); the record's summary repeats on each row, so the
# columns stay fixed whatever KPI set a configuration uses.
CSV_COLUMNS = (
    "Employee ID", "Employee Name", "Department", "Year", "Month", "Status", "Overall Score",
    "Performance Band", "Manager Remark", "KPI Code", "KPI Name", "Weight", "Score", "Source",
    "KPI Remarks",
)
NUMBER_FORMAT = "0.00"
XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def employee_id(employee) -> str:
    """The business employee code; the internal id when no code is set."""
    return employee.employee_code or str(employee.pk)


def _text(value) -> str:
    return "" if value is None else str(value)


def summary_row(record) -> tuple:
    employee = record.employee
    return (
        employee_id(employee), employee.full_name, employee.department.code, record.year,
        record.month, record.status, record.overall_score, record.performance_band,
    )


def kpi_rows(record):
    employee = record.employee
    for score in record.report_scores:
        yield (
            employee_id(employee), employee.full_name, record.year, record.month,
            score.kpi.code, score.kpi.name, score.weight, score.score, score.source,
            score.manager_remark,
        )


def write_csv(records) -> bytes:
    """UTF-8 with a byte-order mark (so Excel shows non-ASCII names correctly); standard CSV
    quoting; one header row even when there is no data."""
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\r\n")
    writer.writerow(CSV_COLUMNS)
    for record in records:
        summary = summary_row(record)
        for kpi in kpi_rows(record):
            writer.writerow([
                *(_text(v) for v in summary), record.manager_remark,
                *(_text(v) for v in kpi[4:]),
            ])
    return buffer.getvalue().encode("utf-8-sig")


def _cell(value):
    """Excel cell value: numbers as numbers, blank text as a truly empty cell."""
    if isinstance(value, Decimal):
        return float(value)
    return None if value == "" else value


def _sheet(workbook, title, columns, rows, numeric_columns, first=False):
    sheet = workbook.active if first else workbook.create_sheet()
    sheet.title = title
    sheet.append(list(columns))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    count = 0
    for row in rows:
        sheet.append([_cell(v) for v in row])
        count += 1
    for index, name in enumerate(columns, start=1):
        letter = get_column_letter(index)
        sheet.column_dimensions[letter].width = max(12, len(name) + 4)
        if index in numeric_columns:
            for (cell,) in sheet.iter_rows(min_row=2, min_col=index, max_col=index):
                cell.number_format = NUMBER_FORMAT
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{count + 1}"
    return sheet


def write_xlsx(records) -> bytes:
    records = list(records)
    workbook = Workbook()
    _sheet(workbook, "Performance Summary", SUMMARY_COLUMNS,
           (summary_row(r) for r in records), numeric_columns={7}, first=True)
    _sheet(workbook, "KPI Details", KPI_COLUMNS,
           (row for r in records for row in kpi_rows(r)), numeric_columns={7, 8})
    workbook["Performance Summary"].column_dimensions["B"].width = 28
    workbook["KPI Details"].column_dimensions["B"].width = 28
    workbook["KPI Details"].column_dimensions["J"].width = 40
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
